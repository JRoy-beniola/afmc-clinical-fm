import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document

from afmc_fm.reproducibility.bootstrap import bootstrap_phase
from afmc_fm.reproducibility.models import PhaseDefinition
from afmc_fm.reproducibility.report_manifest import load_manifest


def _phase() -> PhaseDefinition:
    return PhaseDefinition(
        phase_id="fixture",
        official_evidence_root=Path("docs/results/fixture"),
        official_report=Path("docs/results/fixture/report.docx"),
        decision_record=Path("docs/results/fixture/decision.md"),
        expected_classification="FIXTURE STOP",
        result_kind="historical",
        implementation_sha="a" * 40,
        execution_sha="a" * 40,
        protocol_paths=(),
        manifests=(),
        raw_evidence_paths=(),
        derived_table_paths=(),
        figure_paths=(),
        report_source=None,
        environment_status="unknown",
        rebuild_supported=False,
        rerun_supported=False,
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_report(root: Path) -> Path:
    path = root / "docs/results/fixture/report.docx"
    path.parent.mkdir(parents=True)
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Fixture</w:t></w:r></w:p>
    <w:p><w:r><w:t>Body text</w:t></w:r></w:p>
    <w:sectPr/>
  </w:body>
</w:document>
"""
    styles = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="Heading 1"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/styles.xml", styles)
    return path


def test_bootstrap_refuses_historical_destination(tmp_path: Path):
    with pytest.raises(ValueError, match="historical evidence"):
        bootstrap_phase(
            root=tmp_path,
            phase=_phase(),
            destination=tmp_path / "docs/results/phase0/reproducibility",
        )


def test_bootstrap_refuses_reproduction_output_destination(tmp_path: Path):
    with pytest.raises(ValueError, match="reproduction output"):
        bootstrap_phase(
            root=tmp_path,
            phase=_phase(),
            destination=tmp_path / "outputs/reproduction/fixture/bootstrap",
        )


def test_bootstrap_is_deterministic_and_does_not_mutate_source(tmp_path: Path):
    report = _write_report(tmp_path)
    before = _sha256(report)
    first = tmp_path / "staging-one"
    second = tmp_path / "staging-two"

    bootstrap_phase(root=tmp_path, phase=_phase(), destination=first)
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=second)

    expected = {"report-source.md", "report.yaml", "extraction.json"}
    assert {path.name for path in first.iterdir()} == expected
    assert {path.name for path in second.iterdir()} == expected
    for name in expected:
        assert (first / name).read_bytes() == (second / name).read_bytes()
    assert _sha256(report) == before

    manifest = load_manifest(first / "report.yaml")
    assert manifest.report_source == Path("docs/reproducibility/fixture/report-source.md")
    assert manifest.reference_report_sha256 == before
    assert manifest.audit_status == "pending"


def test_bootstrap_preserves_python_docx_heading_semantics(tmp_path: Path):
    report = tmp_path / "docs/results/fixture/report.docx"
    report.parent.mkdir(parents=True)
    document = Document()
    document.add_heading("Fixture heading", level=1)
    document.add_paragraph("Body text")
    document.save(report)

    destination = tmp_path / "staging"
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=destination)

    source = (destination / "report-source.md").read_text(encoding="utf-8")
    assert source.startswith("# Fixture heading\n"), source[:80]


def test_bootstrap_serializes_whitespace_only_paragraph_as_blank_marker(tmp_path: Path):
    report = tmp_path / "docs/results/fixture/report.docx"
    report.parent.mkdir(parents=True)
    document = Document()
    document.add_paragraph("Before")
    whitespace = document.add_paragraph()
    whitespace.add_run().add_break()
    document.add_paragraph("After")
    document.save(report)

    destination = tmp_path / "staging"
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=destination)

    source = (destination / "report-source.md").read_text(encoding="utf-8")
    assert source == "Before\n\n<!-- blank -->\n\nAfter\n"


def test_bootstrap_refuses_overwrite_without_replace(tmp_path: Path):
    _write_report(tmp_path)
    destination = tmp_path / "staging"
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=destination)

    with pytest.raises(FileExistsError, match="already exists"):
        bootstrap_phase(root=tmp_path, phase=_phase(), destination=destination)


def test_replace_is_allowed_only_beneath_docs_reproducibility(tmp_path: Path):
    _write_report(tmp_path)
    staging = tmp_path / "staging"
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=staging)
    with pytest.raises(ValueError, match="replace"):
        bootstrap_phase(
            root=tmp_path,
            phase=_phase(),
            destination=staging,
            replace=True,
        )

    canonical = tmp_path / "docs/reproducibility/fixture"
    bootstrap_phase(root=tmp_path, phase=_phase(), destination=canonical)
    before = {path.name: path.read_bytes() for path in canonical.iterdir()}
    bootstrap_phase(
        root=tmp_path,
        phase=_phase(),
        destination=canonical,
        replace=True,
    )
    after = {path.name: path.read_bytes() for path in canonical.iterdir()}
    assert after == before
