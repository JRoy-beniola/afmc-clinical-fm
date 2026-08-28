import hashlib
from pathlib import Path

import pytest
from docx import Document

from afmc_fm.reproducibility.documents import (
    BlankNode,
    HeadingNode,
    ImagePlaceholderNode,
    ParagraphNode,
    TableNode,
    build_report_docx,
    parse_report_source,
)
from afmc_fm.reproducibility.report_source import extract_docx


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _reference_docx(path: Path) -> None:
    path.parent.mkdir(parents=True)
    document = Document()
    document.add_heading("Reference", level=1)
    document.add_paragraph("Template body")
    document.save(path)


def test_parse_report_source_preserves_supported_block_order():
    source = """# Report title

Narrative line one.
Narrative line two.

<!-- blank -->

| A | B |
| --- | --- |
| 1 | two\\|parts |

<!-- image rel=rId7 -->
"""

    nodes = parse_report_source(source)

    assert nodes == (
        HeadingNode(level=1, text="Report title"),
        ParagraphNode(text="Narrative line one.\nNarrative line two."),
        BlankNode(),
        TableNode(rows=(("A", "B"), ("1", "two|parts"))),
        ImagePlaceholderNode(relationship_id="rId7"),
    )


def test_build_report_docx_is_deterministic_and_preserves_inputs(tmp_path: Path):
    reference = tmp_path / "docs/results/fixture/reference.docx"
    _reference_docx(reference)
    source = tmp_path / "docs/reproducibility/fixture/report-source.md"
    source.parent.mkdir(parents=True)
    source.write_text(
        """# Rebuilt report

Recovered narrative.

| Metric | Value |
| --- | --- |
| MAE | 1.25 |

<!-- image rel=rId9 -->
""",
        encoding="utf-8",
    )
    reference_before = _sha256(reference)
    source_before = _sha256(source)
    first = tmp_path / "outputs/reproduction/fixture/rebuild/one/report.docx"
    second = tmp_path / "outputs/reproduction/fixture/rebuild/two/report.docx"

    build_report_docx(source=source, reference=reference, output=first)
    build_report_docx(source=source, reference=reference, output=second)

    assert _sha256(first) == _sha256(second)
    assert _sha256(reference) == reference_before
    assert _sha256(source) == source_before
    snapshot = extract_docx(first)
    paragraph_text = [block.text for block in snapshot.blocks if hasattr(block, "text")]
    assert "Rebuilt report" in paragraph_text
    assert "Recovered narrative." in paragraph_text
    assert "[image rel=rId9]" in paragraph_text
    assert snapshot.table_count == 1


def test_build_report_docx_refuses_historical_output(tmp_path: Path):
    reference = tmp_path / "docs/results/fixture/reference.docx"
    _reference_docx(reference)
    source = tmp_path / "docs/reproducibility/fixture/report-source.md"
    source.parent.mkdir(parents=True)
    source.write_text("# Report\n", encoding="utf-8")

    with pytest.raises(ValueError, match="historical"):
        build_report_docx(
            source=source,
            reference=reference,
            output=tmp_path / "docs/results/fixture/rebuilt.docx",
        )
