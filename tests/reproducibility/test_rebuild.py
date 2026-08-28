import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from afmc_fm.reproducibility.models import PhaseDefinition
from afmc_fm.reproducibility.rebuild import rebuild_phase
from afmc_fm.reproducibility.registry import PHASES, get_phase
from afmc_fm.reproducibility.report_manifest import build_manifest, dump_manifest
from afmc_fm.reproducibility.report_source import extract_docx, render_markdown


def _tree_snapshot(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def _fixture_phase(tmp_path: Path, monkeypatch) -> PhaseDefinition:
    archive = tmp_path / "docs/results/fixture"
    archive.mkdir(parents=True)
    decision = archive / "decision.md"
    decision.write_text("FIXTURE COMPLETE\n", encoding="utf-8")
    reference = archive / "report.docx"
    document = Document()
    document.add_heading("Fixture report", level=1)
    document.add_paragraph("Recovered narrative.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "MAE"
    table.cell(1, 1).text = "1.25"
    document.save(reference)
    legacy = archive / "legacy.csv"
    legacy.write_text("metric,value\nMAE,1.25\n", encoding="utf-8")

    source_dir = tmp_path / "docs/reproducibility/fixture"
    source_dir.mkdir(parents=True)
    snapshot = extract_docx(reference)
    source_bytes = render_markdown(snapshot).encode("utf-8")
    source = source_dir / "report-source.md"
    source.write_bytes(source_bytes)

    phase = PhaseDefinition(
        phase_id="fixture",
        official_evidence_root=Path("docs/results/fixture"),
        official_report=Path("docs/results/fixture/report.docx"),
        decision_record=Path("docs/results/fixture/decision.md"),
        expected_classification="FIXTURE COMPLETE",
        result_kind="historical",
        implementation_sha="a" * 40,
        execution_sha=None,
        protocol_paths=(),
        manifests=(),
        raw_evidence_paths=(),
        derived_table_paths=(Path("docs/results/fixture/legacy.csv"),),
        figure_paths=(),
        report_source=Path("docs/reproducibility/fixture/report-source.md"),
        environment_status="unknown",
        rebuild_supported=True,
        rerun_supported=False,
    )
    monkeypatch.setitem(PHASES, "fixture", phase)

    manifest = build_manifest(
        phase,
        snapshot,
        Path("docs/reproducibility/fixture/report-source.md"),
        source_bytes,
        audit_status="accepted",
    )
    (source_dir / "report.yaml").write_text(dump_manifest(manifest), encoding="utf-8")
    (source_dir / "extraction.json").write_text(
        json.dumps(snapshot.to_json_dict(), ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    (source_dir / "rebuild.yaml").write_text(
        """schema_version: 1
phase_id: fixture
report_source: docs/reproducibility/fixture/report-source.md
reference_report: docs/results/fixture/report.docx
document_output: report/candidate.docx
artifacts:
  - kind: legacy-table
    source: docs/results/fixture/legacy.csv
    output: reference/legacy.csv
    mode: reference-copy
""",
        encoding="utf-8",
    )
    return phase


def test_rebuild_phase_orchestrates_without_mutating_official_archive(tmp_path: Path, monkeypatch):
    phase = _fixture_phase(tmp_path, monkeypatch)
    before = _tree_snapshot(tmp_path / phase.official_evidence_root)

    report = rebuild_phase(tmp_path, "fixture")

    after = _tree_snapshot(tmp_path / phase.official_evidence_root)
    expected_root = (tmp_path / "outputs/reproduction/fixture/rebuild").resolve()
    assert report.structural_comparison.ok, report.structural_comparison.to_json_dict()
    assert report.ok
    assert report.destination == expected_root
    assert report.candidate_report == expected_root / "report/candidate.docx"
    assert report.candidate_report.is_file()
    assert report.rebuild_report == expected_root / "rebuild-report.json"
    assert report.rebuild_report.is_file()
    assert report.reference_copy_count == 1
    assert report.generated_count == 0
    assert report.historical_archive_unchanged
    assert before == after

    payload = json.loads(report.rebuild_report.read_text(encoding="utf-8"))
    assert payload["phase_id"] == "fixture"
    assert payload["historical_archive_unchanged"] is True
    assert payload["artifact_counts"] == {"generated": 0, "reference_copy": 1}
    assert payload["structural_comparison"]["byte_identical"] is False


def test_rebuild_phase_rejects_unaccepted_report_source(tmp_path: Path, monkeypatch):
    _fixture_phase(tmp_path, monkeypatch)
    manifest_path = tmp_path / "docs/reproducibility/fixture/report.yaml"
    text = manifest_path.read_text(encoding="utf-8").replace(
        "audit_status: accepted",
        "audit_status: pending",
    )
    manifest_path.write_text(text, encoding="utf-8")

    try:
        rebuild_phase(tmp_path, "fixture")
    except ValueError as exc:
        assert "accepted" in str(exc)
    else:
        raise AssertionError("pending R2 source manifest must not be rebuildable")


@pytest.mark.parametrize("phase_id", ["phase0", "phase05", "phase06"])
def test_registered_historical_phase_rebuilds_without_archive_mutation(phase_id: str):
    repository = Path.cwd()
    phase = get_phase(phase_id)
    before = _tree_snapshot(repository / phase.official_evidence_root)

    report = rebuild_phase(repository, phase_id)

    after = _tree_snapshot(repository / phase.official_evidence_root)
    assert report.structural_comparison.ok, report.structural_comparison.to_json_dict()
    assert report.ok
    assert report.destination == repository / "outputs/reproduction" / phase_id / "rebuild"
    assert report.candidate_report.is_file()
    assert report.rebuild_report.is_file()
    assert report.generated_count == 0
    assert report.reference_copy_count == len(report.artifacts)
    assert all(result.mode == "reference-copy" for result in report.artifacts)
    assert report.historical_archive_unchanged
    assert before == after
