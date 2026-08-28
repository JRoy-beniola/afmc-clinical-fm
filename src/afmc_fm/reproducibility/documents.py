from __future__ import annotations

import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from docx import Document

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
_IMAGE_RE = re.compile(r"^<!-- image rel=([^ ]+) -->$")
_TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_CORE_TIME = datetime(2000, 1, 1, 0, 0, 0, tzinfo=UTC)


@dataclass(frozen=True)
class HeadingNode:
    level: int
    text: str


@dataclass(frozen=True)
class ParagraphNode:
    text: str


@dataclass(frozen=True)
class BlankNode:
    pass


@dataclass(frozen=True)
class TableNode:
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class ImagePlaceholderNode:
    relationship_id: str


DocumentNode = HeadingNode | ParagraphNode | BlankNode | TableNode | ImagePlaceholderNode


def _split_table_row(line: str) -> tuple[str, ...]:
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        raise ValueError("table row must start and end with '|'")

    cells: list[str] = []
    current: list[str] = []
    index = 1
    end = len(stripped) - 1
    while index < end:
        char = stripped[index]
        if char == "\\" and index + 1 < end:
            current.append(stripped[index + 1])
            index += 2
            continue
        if char == "|":
            cells.append("".join(current).strip().replace("<br>", "\n"))
            current = []
        else:
            current.append(char)
        index += 1
    cells.append("".join(current).strip().replace("<br>", "\n"))
    return tuple(cells)


def _table_at(lines: list[str], index: int) -> tuple[TableNode, int] | None:
    if not lines[index].lstrip().startswith("|"):
        return None
    table_lines: list[str] = []
    cursor = index
    while cursor < len(lines) and lines[cursor].lstrip().startswith("|"):
        table_lines.append(lines[cursor])
        cursor += 1
    if len(table_lines) < 2:
        return None
    header = _split_table_row(table_lines[0])
    separator = _split_table_row(table_lines[1])
    if len(header) != len(separator) or not all(
        _TABLE_SEPARATOR_RE.fullmatch(cell.strip()) for cell in separator
    ):
        return None
    rows = [header]
    for line in table_lines[2:]:
        row = _split_table_row(line)
        if len(row) != len(header):
            raise ValueError("table row width does not match header width")
        rows.append(row)
    return TableNode(rows=tuple(rows)), cursor


def parse_report_source(source: str) -> tuple[DocumentNode, ...]:
    """Parse the deterministic R2 Markdown representation into document nodes."""

    lines = source.splitlines()
    nodes: list[DocumentNode] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.strip() == "<!-- blank -->":
            nodes.append(BlankNode())
            index += 1
            continue
        image = _IMAGE_RE.fullmatch(line.strip())
        if image:
            nodes.append(ImagePlaceholderNode(relationship_id=image.group(1)))
            index += 1
            continue
        heading = _HEADING_RE.fullmatch(line)
        if heading:
            nodes.append(HeadingNode(level=len(heading.group(1)), text=heading.group(2)))
            index += 1
            continue
        table = _table_at(lines, index)
        if table is not None:
            node, index = table
            nodes.append(node)
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines) and lines[index].strip():
            candidate = lines[index]
            if (
                candidate.strip() == "<!-- blank -->"
                or _IMAGE_RE.fullmatch(candidate.strip())
                or _HEADING_RE.fullmatch(candidate)
                or _table_at(lines, index) is not None
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        nodes.append(ParagraphNode(text="\n".join(paragraph_lines)))
    return tuple(nodes)


def _is_historical(path: Path) -> bool:
    parts = path.resolve().parts
    return any(parts[index : index + 2] == ("docs", "results") for index in range(len(parts) - 1))


def _clear_body(document: Document) -> None:
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def _populate(document: Document, nodes: tuple[DocumentNode, ...]) -> None:
    for node in nodes:
        if isinstance(node, HeadingNode):
            document.add_heading(node.text, level=node.level)
        elif isinstance(node, ParagraphNode):
            document.add_paragraph(node.text)
        elif isinstance(node, BlankNode):
            document.add_paragraph("")
        elif isinstance(node, TableNode):
            width = max((len(row) for row in node.rows), default=1)
            table = document.add_table(rows=len(node.rows), cols=width)
            for row_index, row in enumerate(node.rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
        else:
            document.add_paragraph(f"[image rel={node.relationship_id}]")


def _neutralize_properties(document: Document) -> None:
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = _FIXED_CORE_TIME
    properties.modified = _FIXED_CORE_TIME
    properties.last_printed = None
    properties.revision = 1


def _deterministic_repack(source: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp")
    try:
        with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
            temporary,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as rebuilt:
            for name in sorted(archive.namelist()):
                info = zipfile.ZipInfo(name, _FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0
                rebuilt.writestr(info, archive.read(name), compress_type=zipfile.ZIP_DEFLATED)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()


def build_report_docx(*, source: Path, reference: Path, output: Path) -> Path:
    """Build a deterministic DOCX from recovered source and a read-only style reference."""

    source_path = Path(source).resolve()
    reference_path = Path(reference).resolve()
    output_path = Path(output).resolve()
    if _is_historical(output_path):
        raise ValueError("rebuilt report output may not be historical evidence")
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    if not reference_path.is_file():
        raise FileNotFoundError(reference_path)

    nodes = parse_report_source(source_path.read_text(encoding="utf-8"))
    document = Document(reference_path)
    _clear_body(document)
    _populate(document, nodes)
    _neutralize_properties(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="afmc-rebuild-") as directory:
        raw = Path(directory) / "report.docx"
        document.save(raw)
        _deterministic_repack(raw, output_path)
    return output_path
